import re
import openpyxl
from datetime import datetime
from django.db import transaction
from django.db.models import Q

from .models import Episode, EpisodeShow, Transcript, Character, TranscriptLanguage


def get_episodes_for_user(user):
    from apps.shows.permissions import get_visible_show_ids_for_user
    visible_ids = get_visible_show_ids_for_user(user)
    return Episode.objects.filter(
        Q(primary_show_id__in=visible_ids) | Q(shows__id__in=visible_ids)
    ).distinct().select_related('primary_show', 'summary').prefetch_related(
        'transcripts', 'characters', 'episodeshow_set__show'
    )


def add_episode_to_show(episode: Episode, show_id: str, order: int = None):
    EpisodeShow.objects.get_or_create(
        episode=episode, show_id=show_id,
        defaults={'episode_order': order},
    )


def remove_episode_from_show(episode: Episode, show_id: str):
    EpisodeShow.objects.filter(episode=episode, show_id=show_id).delete()


def _parse_flexible_date(date_str):
    if not date_str:
        return None
    cleaned = re.sub(r'(\d+)(st|nd|rd|th)', r'\1', str(date_str).strip(), flags=re.IGNORECASE)
    for fmt in ('%B %d, %Y', '%B %d %Y', '%Y-%m-%d', '%d/%m/%Y', '%m/%d/%Y', '%d-%m-%Y'):
        try:
            return datetime.strptime(cleaned, fmt).date()
        except ValueError:
            continue
    return None


def _get_or_create_character(character_ref: str, name: str, actor: str) -> Character:
    char, _ = Character.objects.get_or_create(
        character_ref=character_ref.strip(),
        actor=(actor or '').strip(),
        defaults={'name': name.strip()},
    )
    return char


def _parse_dialogue_text(text: str, character_map: dict) -> list:
    """Split a multi-line dialogue block into structured rows."""
    rows = []
    if not text:
        return rows
    for line in str(text).split('\n'):
        line = line.strip()
        if not line:
            continue
        match = re.match(r'^(.+?):\s*["\']?(.+?)["\']?\s*$', line)
        if match:
            char_name = match.group(1).strip()
            dialogue = match.group(2).strip()
            char_ref = next(
                (ref for ref, info in character_map.items()
                 if info['name'].lower() == char_name.lower()),
                char_name.lower().replace(' ', '_'),
            )
            rows.append({'character_ref': char_ref, 'character_name': char_name, 'text': dialogue})
        elif line:
            rows.append({'character_ref': '_narrator', 'character_name': 'Narrator', 'text': line})
    return rows


@transaction.atomic
def parse_episode_excel(episode: Episode, file) -> Transcript:
    """Parse per-episode Excel (scene_id, timecode, character, text, notes) → origin Transcript."""
    wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
    ws = wb.active
    rows_data = []
    headers = None
    characters_seen: dict[str, str] = {}

    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            headers = [str(c).strip().lower() if c else f'col_{j}' for j, c in enumerate(row)]
            continue
        if not any(row):
            continue
        row_dict = dict(zip(headers, row))
        char_name = str(row_dict.get('character', '')).strip()
        text = str(row_dict.get('text', '')).strip()
        if not text:
            continue
        char_ref = char_name.lower().replace(' ', '_')
        characters_seen[char_ref] = char_name
        rows_data.append({'character_ref': char_ref, 'character_name': char_name, 'text': text})

    wb.close()

    ep_chars = [_get_or_create_character(ref, name, '') for ref, name in characters_seen.items()]
    if ep_chars:
        episode.characters.add(*ep_chars)

    transcript, _ = Transcript.objects.update_or_create(
        episode=episode, language=TranscriptLanguage.ORIGIN,
        defaults={'rows': rows_data},
    )
    return transcript


@transaction.atomic
def parse_show_excel(file, show, root) -> list:
    """Parse a whole-node Excel upload → Episodes with origin Transcripts. `show` is the
    node the file was uploaded to (any depth); `root` is its top-level ancestor, used as
    each episode's primary_show for display (matches the pre-existing convention that
    primary_show always points at a root show). Returns created episodes."""
    wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
    ws = wb.active

    episode_groups: dict[str, list] = {}
    headers = None

    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            headers = [str(c).strip().lower() if c else f'col_{j}' for j, c in enumerate(row)]
            continue
        if not any(row):
            continue
        row_dict = dict(zip(headers, row))
        ep_id = str(row_dict.get('episode_id', '')).strip()
        if not ep_id:
            continue
        if ep_id not in episode_groups:
            episode_groups[ep_id] = []
        episode_groups[ep_id].append(row_dict)

    wb.close()

    created_episodes = []

    for ep_order, (ep_id, scene_rows) in enumerate(episode_groups.items()):
        # Build character map: ref → {name, actor}
        character_map: dict[str, dict] = {}
        for scene in scene_rows:
            for prefix in ('character_a', 'character_b'):
                ref = str(scene.get(f'{prefix}_id', '')).strip()
                name = str(scene.get(f'{prefix}_name', '')).strip()
                actor = str(scene.get(f'{prefix}_actor', '')).strip()
                if ref and name:
                    character_map[ref] = {'name': name, 'actor': actor}

        # Resolve air_date from any column whose header contains "air" or "date"
        air_date = None
        first_scene = scene_rows[0]
        for key in (headers or []):
            if ('air' in key or 'date' in key) and key != 'updated_at':
                air_date = _parse_flexible_date(first_scene.get(key))
                if air_date:
                    break

        episode, _ = Episode.objects.update_or_create(
            primary_show=root,
            episode_number=ep_id,
            defaults={'title': ep_id, 'air_date': air_date},
        )
        add_episode_to_show(episode, str(show.id), order=ep_order)

        ep_chars = [
            _get_or_create_character(ref, info['name'], info['actor'])
            for ref, info in character_map.items()
        ]
        if ep_chars:
            episode.characters.add(*ep_chars)

        all_rows = []
        for scene in scene_rows:
            all_rows.extend(_parse_dialogue_text(scene.get('text', ''), character_map))

        Transcript.objects.update_or_create(
            episode=episode, language=TranscriptLanguage.ORIGIN,
            defaults={'rows': all_rows},
        )
        created_episodes.append(episode)

    return created_episodes
